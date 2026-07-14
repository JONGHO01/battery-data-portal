import math
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt


NS_RE = re.compile(r"^\{.*\}")
NUM_RE = re.compile(r"[-+]?(?:\d*\.\d+|\d+)(?:[eE][-+]?\d+)?")
CMD_RE = re.compile(r"([ML])\s*([-+]?(?:\d*\.\d+|\d+)(?:[eE][-+]?\d+)?)\s*,?\s*([-+]?(?:\d*\.\d+|\d+)(?:[eE][-+]?\d+)?)")


def local_name(tag):
    return NS_RE.sub("", tag)


def f(value, default=0.0):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def color(value, fallback="111827"):
    value = (value or fallback).strip()
    if value.startswith("#"):
        value = value[1:]
    if len(value) == 3:
        value = "".join(ch * 2 for ch in value)
    try:
        return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))
    except Exception:
        return RGBColor(int(fallback[0:2], 16), int(fallback[2:4], 16), int(fallback[4:6], 16))


def class_set(el):
    return set((el.attrib.get("class") or "").split())


def parse_viewbox(root):
    raw = root.attrib.get("viewBox", "")
    nums = [float(x) for x in NUM_RE.findall(raw)]
    if len(nums) == 4:
        return nums
    return [0.0, 0.0, f(root.attrib.get("width"), 1120), f(root.attrib.get("height"), 720)]


def parse_transform(raw):
    tx = ty = rot = 0.0
    for name, args in re.findall(r"(translate|rotate)\(([^)]*)\)", raw or ""):
        nums = [float(x) for x in NUM_RE.findall(args)]
        if name == "translate" and nums:
            tx += nums[0]
            ty += nums[1] if len(nums) > 1 else 0.0
        elif name == "rotate" and nums:
            rot += nums[0]
    return tx, ty, rot


class SvgToPpt:
    def __init__(self, slide, root, slide_w, slide_h):
        self.slide = slide
        self.vx, self.vy, self.vw, self.vh = parse_viewbox(root)
        margin_x = Inches(0.55)
        margin_y = Inches(0.38)
        self.scale = min((slide_w - margin_x * 2) / self.vw, (slide_h - margin_y * 2) / self.vh)
        self.left = (slide_w - self.vw * self.scale) / 2
        self.top = (slide_h - self.vh * self.scale) / 2

    def x(self, value):
        return int(self.left + (value - self.vx) * self.scale)

    def y(self, value):
        return int(self.top + (value - self.vy) * self.scale)

    def length(self, value):
        return int(value * self.scale)

    def pt(self, x, y):
        return self.x(x), self.y(y)

    def line_style(self, shape, el):
        cls = class_set(el)
        stroke = el.attrib.get("stroke")
        width = f(el.attrib.get("stroke-width"), 1.0)
        if "rateLine" in cls:
            stroke = "8A8F98"
            width = 0.9
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        elif "grid" in cls:
            stroke = "E5E7EB"
            width = 0.7
        elif "axis" in cls or "frame" in cls:
            stroke = "111827"
            width = 1.1
        elif "capLine" in cls:
            stroke = "EC7DBF"
            width = 1.2
        shape.line.color.rgb = color(stroke, "111827")
        shape.line.width = Pt(width)

    def add_line(self, el, tx, ty):
        x1, y1 = self.pt(f(el.attrib.get("x1")) + tx, f(el.attrib.get("y1")) + ty)
        x2, y2 = self.pt(f(el.attrib.get("x2")) + tx, f(el.attrib.get("y2")) + ty)
        shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
        self.line_style(shape, el)

    def add_rect(self, el, tx, ty):
        cls = class_set(el)
        w = f(el.attrib.get("width"))
        h = f(el.attrib.get("height"))
        if "frame" not in cls:
            return
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            self.x(f(el.attrib.get("x")) + tx),
            self.y(f(el.attrib.get("y")) + ty),
            self.length(w),
            self.length(h),
        )
        shape.fill.background()
        self.line_style(shape, el)

    def add_circle(self, el, tx, ty):
        cls = class_set(el)
        r = f(el.attrib.get("r"), 3.0)
        cx = f(el.attrib.get("cx")) + tx
        cy = f(el.attrib.get("cy")) + ty
        left = self.x(cx - r)
        top = self.y(cy - r)
        shape = self.slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, self.length(r * 2), self.length(r * 2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.line.color.rgb = color("EC7DBF" if "capPt" in cls else "0F766E")
        shape.line.width = Pt(1.35 if "capPt" in cls else 1.0)

    def add_text(self, el, tx, ty, rot):
        text = "".join(el.itertext()).strip()
        if not text:
            return
        cls = class_set(el)
        size = 9
        if "label" in cls:
            size = 14
        elif "rateLabel" in cls:
            size = 11
        elif "legend" in cls or "caption" in cls:
            size = 8
        elif "panelTitle" in cls:
            size = 15
        elif "tick" in cls:
            size = 9

        x = f(el.attrib.get("x")) + tx
        y = f(el.attrib.get("y")) + ty
        est_w = max(28, len(text) * size * 0.62)
        est_h = size * 1.45
        anchor = el.attrib.get("text-anchor", "")
        left_svg = x - est_w / 2 if anchor == "middle" else x - est_w if anchor == "end" else x
        top_svg = y - est_h * 0.8
        box = self.slide.shapes.add_textbox(self.x(left_svg), self.y(top_svg), self.length(est_w), self.length(est_h))
        if rot:
            box.rotation = rot
        tf = box.text_frame
        tf.clear()
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = text
        run.font.name = "Segoe UI"
        run.font.size = Pt(size)
        run.font.bold = "panelTitle" in cls
        run.font.color.rgb = color("111827" if "label" in cls or "rateLabel" in cls else "374151")

    def add_path(self, el, tx, ty):
        cls = class_set(el)
        if "capLine" not in cls:
            return
        points = []
        for cmd, xs, ys in CMD_RE.findall(el.attrib.get("d", "")):
            points.append((cmd, float(xs) + tx, float(ys) + ty))
        prev = None
        for cmd, x, y in points:
            if cmd == "M" or prev is None:
                prev = (x, y)
                continue
            x1, y1 = self.pt(prev[0], prev[1])
            x2, y2 = self.pt(x, y)
            shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
            self.line_style(shape, el)
            prev = (x, y)

    def walk(self, el, tx=0.0, ty=0.0, rot=0.0):
        add_tx, add_ty, add_rot = parse_transform(el.attrib.get("transform", ""))
        tx += add_tx
        ty += add_ty
        rot += add_rot
        tag = local_name(el.tag)
        if tag == "line":
            self.add_line(el, tx, ty)
        elif tag == "rect":
            self.add_rect(el, tx, ty)
        elif tag == "circle":
            self.add_circle(el, tx, ty)
        elif tag == "path":
            self.add_path(el, tx, ty)
        elif tag == "text":
            self.add_text(el, tx, ty, rot)
            return
        for child in list(el):
            if local_name(child.tag) != "style":
                self.walk(child, tx, ty, rot)


def add_png_slide(slide, image_path, slide_w, slide_h):
    margin_x = Inches(0.55)
    margin_y = Inches(0.45)
    max_w = slide_w - margin_x * 2
    max_h = slide_h - margin_y * 2
    with Image.open(image_path) as img:
        img_w, img_h = img.size
    scale = min(max_w / img_w, max_h / img_h)
    pic_w = int(img_w * scale)
    pic_h = int(img_h * scale)
    slide.shapes.add_picture(str(image_path), int((slide_w - pic_w) / 2), int((slide_h - pic_h) / 2), width=pic_w, height=pic_h)


def main():
    if len(sys.argv) < 3:
        raise SystemExit("usage: export_ppt.py input.png output.pptx [input.svg]")

    image_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    svg_path = Path(sys.argv[3]) if len(sys.argv) > 3 else None

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    rendered = False
    if svg_path and svg_path.exists() and svg_path.stat().st_size:
        try:
            root = ET.parse(svg_path).getroot()
            SvgToPpt(slide, root, prs.slide_width, prs.slide_height).walk(root)
            rendered = True
        except Exception:
            rendered = False

    if not rendered:
        add_png_slide(slide, image_path, prs.slide_width, prs.slide_height)

    prs.save(output_path)


if __name__ == "__main__":
    main()
